import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_final_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean Modern Academic Palette
    DARK_BG = RGBColor(17, 24, 39)        # Clean Charcoal Slate
    CARD_BG = RGBColor(31, 41, 55)        # Card Slate
    ACCENT_BLUE = RGBColor(59, 130, 246)  # Electric Blue
    ACCENT_TEAL = RGBColor(20, 184, 166)  # Teal
    TEXT_WHITE = RGBColor(243, 244, 246)  # Off-White
    TEXT_MUTED = RGBColor(156, 163, 175)  # Muted Gray

    slides_data = [
        {
            "title": "OmniTutor: Source-Grounded Multimodal Study Assistant",
            "subtitle": "Track D: Personalized Tutoring & Adaptive Learning | Multimodal AI Hackathon 2026 (IIT Mandi)",
            "bullets": [
                "Team Name: CODEFORGE",
                "Team Leader: PIRUTHVI M K (7376241EE166) — 3rd Year EEE",
                "Team Members:",
                "  • SRIRAM B (7376241EE195) — 3rd Year EEE",
                "  • VIKAS P (7376241EE212) — 3rd Year EEE",
                "  • GOBUNATH D (7376241EE125) — 3rd Year EEE",
                "Faculty Mentor: Prof. SRITHA P (Staff ID: EE1747) — Dept. of EEE",
                "Institution: Bannari Amman Institute of Technology, Sathyamangalam"
            ],
            "notes": "Good morning judges and mentors. We are Team CODEFORGE from Bannari Amman Institute of Technology, presenting OmniTutor — an AI study companion that indexes lecture videos, slides, and question bank PDFs to deliver source-cited answers with exact timestamps and adaptive practice tests."
        },
        {
            "title": "Why We Built This: Real Student Pain Points",
            "subtitle": "Practical challenges faced by engineering students before exams",
            "bullets": [
                "Time Wasted Finding Explanations: Finding a 2-minute formula derivation in a 90-minute lecture takes 20+ minutes of manual scrubbing.",
                "Disconnected Course Materials: Semester question banks, PPT slides, and recorded lecture videos exist in separate silos.",
                "ChatGPT / Gemini Hallucinations: Standard LLMs provide generic textbook answers that frequently do not match syllabus requirements.",
                "Static Practice Tests: Regular question banks cannot identify root-cause misconceptions or adapt question difficulty to individual student weaknesses."
            ],
            "notes": "As 3rd year engineering students, we constantly face the issue of scattered resources. Generic AI chatbots fail to give exact video references or match the specific course syllabus taught in class."
        },
        {
            "title": "Our Solution: How OmniTutor Works",
            "subtitle": "Turning course materials into an interactive, source-grounded study partner",
            "bullets": [
                "All-in-One Multi-Format Ingestion: Ingests MP4 video lectures, MP3 audio, and course question bank PDFs seamlessly.",
                "Verified Timestamps & Page References: Every explanation provides exact video markers like [Video: 12:40] and slide numbers [Slide: 14].",
                "Adaptive Practice Quizzes: Generates dynamic MCQs by module and automatically scales difficulty (Easy, Medium, Hard).",
                "Targeted Revision Plan: Pinpoints missed questions and highlights the exact video timestamp range and slide page for remedial review."
            ],
            "notes": "OmniTutor unifies lecture recordings and question bank PDFs, answers doubts with verified timestamp citations, and generates personalized quizzes that adapt to your mastery level."
        },
        {
            "title": "System Pipeline & Implementation",
            "subtitle": "How the data flows from raw media to interactive intelligence",
            "bullets": [
                "1. Document Ingestion: Extracts audio transcripts, slide text, and OCR diagrams via PyPDF and video frame parsers.",
                "2. Multimodal RAG Engine: Passes media handles to Gemini 1.5 (1M context) to retrieve answers grounded strictly on uploaded materials.",
                "3. Diagnostic Agent: Analyzes quiz submissions, calculates topic mastery percentages, and maps student weak spots.",
                "4. Interactive Dashboard: Responsive web frontend built with FastAPI, Tailwind CSS, and Chart.js for radar analytics."
            ],
            "notes": "Our pipeline parses lecture media and PDFs, feeds them into a multimodal RAG engine for zero-hallucination citations, and uses an agentic evaluation module to track student mastery."
        },
        {
            "title": "Core Features & Prototype Demo",
            "subtitle": "Tested live with real engineering question banks & lectures",
            "bullets": [
                "Instant Citation Q&A: Ask any conceptual or numerical doubt and get step-by-step math with verified video timecodes.",
                "Topic Extraction: Automatically identifies core syllabus modules from the uploaded PDF/video.",
                "Interactive Quiz Engine: Dynamic MCQs with immediate feedback, option explanations, and source references.",
                "Competency Radar Chart: Visual 360-degree graph showing exact proficiency across each course module."
            ],
            "notes": "We tested OmniTutor with our departmental Question Bank (22EE501). The system accurately retrieved analytical formulas with page citations and updated our competency radar graph in real time."
        },
        {
            "title": "Technical Stack & Reliability",
            "subtitle": "Built for speed, accuracy, and zero downtime",
            "bullets": [
                "AI Models: Google Gemini 1.5 Flash (for high-speed multimodal reasoning) with auto-model fallback.",
                "Backend Framework: Python 3.14 + FastAPI + Uvicorn for lightweight REST API endpoints.",
                "Local Document Fallback: Built-in local PDF parser ensures the app works smoothly even in offline demo mode.",
                "Frontend: Modern HTML5, Tailwind CSS, Lucide icons, and Chart.js visualizations."
            ],
            "notes": "We chose FastAPI and Python for low latency and high reliability. We also built an offline document fallback so the system never crashes during a live demo."
        },
        {
            "title": "Future Scope & Conclusion",
            "subtitle": "Next development milestones for Team CODEFORGE",
            "bullets": [
                "Voice Tutor Mode: Adding speech-to-speech interaction for hands-free problem solving.",
                "College LMS Integration: Connecting directly to Moodle / Google Classroom so professors can see class-wide weak areas.",
                "Exam Readiness Score: Automated score predictor based on practice test consistency.",
                "Summary: OmniTutor saves student study time, eliminates AI hallucinations, and makes exam preparation personalized."
            ],
            "notes": "Our future roadmap includes voice interaction and college LMS integration. Thank you for your time, and Team CODEFORGE is ready for your questions!"
        }
    ]

    blank_layout = prs.slide_layouts[6]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        # Header Box
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE

        tf_card = card.text_frame
        tf_card.word_wrap = True
        p_title = tf_card.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        p_sub = tf_card.add_paragraph()
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = ACCENT_TEAL

        # Content Box
        content_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(5.0))
        content_card.fill.solid()
        content_card.fill.fore_color.rgb = CARD_BG
        content_card.line.fill.background()

        tf_content = content_card.text_frame
        tf_content.word_wrap = True

        for idx, bullet in enumerate(data["bullets"]):
            p = tf_content.paragraphs[0] if idx == 0 else tf_content.add_paragraph()
            p.text = "• " + bullet
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_WHITE
            p.space_after = Pt(8)

        # Notes
        notes_slide = slide.notes_slide
        tf_notes = notes_slide.notes_text_frame
        tf_notes.text = data["notes"]

    output_path = "CODEFORGE_OmniTutor_Final_Submission.pptx"
    prs.save(output_path)
    print(f"Final presentation saved to {output_path}")

if __name__ == "__main__":
    create_final_presentation()
